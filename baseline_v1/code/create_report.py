# -*- coding: utf-8 -*-
import csv, os, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
import model_solver as ms
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parents[1]
DOCS = BASE / 'docs'
RESULTS = BASE / 'results'
FIGS = BASE / 'figures'
DOCS.mkdir(exist_ok=True)

def read_csv(name):
    with open(RESULTS/name, encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

summary = read_csv('summary.csv')
table1 = read_csv('table1_question1.csv')
table2 = read_csv('table2_question2.csv')
table3 = read_csv('table3_question3.csv')
table4 = read_csv('table4_question4.csv')
table5 = read_csv('table5_purchase_plan.csv')
processes = ms.build_processes()

# Markdown report
md = []
md.append('# 2026年五一数学建模B题：多工序协同作业问题研究报告\n')
md.append('## 摘要\n')
md.append('本报告将题目抽象为带车间内工序链约束、设备类型可替代、同设备跨车间转运时间的柔性作业车间调度问题，并建立混合整数线性规划模型。固定设备条件下，Q1、Q2、Q3均由MILP求得全局最优解；Q4通过关键路径下界证明，在题设工序效率不变的前提下，追加设备无法突破C车间流程链下界，因此以“工期优先、成本次优”为准则，最优采购方案为不采购。\n')
md.append('## 主要结果\n')
md.append('| 问题 | 最短时长(s) | 最短时长 | 说明 |\n|---|---:|---|---|')
for r in summary:
    md.append(f"| {r['问题']} | {r['最短时长(s)']} | {r['最短时长(HH:MM:SS)']} | {r['求解状态']} |")
md.append('\n')
md.append('## 建模假设\n')
md.append('1. 设备从对应班组驻地出发，首个车间作业开始时间必须不早于“班组-车间”转运时间。\n2. 同一台设备在同一车间内不同工序间转运时间为0；跨车间转运时间为距离/2m/s，按秒计。\n3. 一个工序若需要两类设备，则两类设备同一时刻开始作业，分别占用其实际工作时长；该工序完工时刻为两类设备结束时刻的最大值。\n4. C3-C5按题意展开为3轮，顺序为C3_R1-C4_R1-C5_R1-C3_R2-C4_R2-C5_R2-C3_R3-C4_R3-C5_R3。\n')
md.append('## 数学模型\n')
md.append('设工序集合为J，设备子任务集合为O，设备集合为M。工序j的开始时刻为S_j，最大完工时间为T。子任务o属于工序j(o)，需设备类型tau(o)，处理时长p_o。若设备m可处理o，则x_{om}=1表示o分配给m；若同一设备m上子任务a排在b之前，则y_{abm}=1。目标为min T。核心约束包括：工序链约束、子任务唯一设备分配、同机不重叠约束、序列相关转运时间约束、首站到达约束和最大完工时间约束。\n')
md.append('## Q4关键路径证明\n')
md.append('C车间链路必须严格顺序完成。其不含初始转运的工序链最短时间为：C1 10368s + C2 7406s + 3*(C3 6480s + C4 14400s + C5 14400s) = 123614s。两班组中到C车间最快为班组1，460m/2m/s=230s，故全局工期下界为123844s。Q3调度已达到123844s，因此Q4即使追加设备也无法缩短总工期。\n')
md.append('## 交付文件说明\n')
md.append('- `results/table1_question1.csv` 到 `results/table5_purchase_plan.csv`：五个问题的可提交结果表。\n- `code/model_solver.py`：完整MILP建模与求解代码。\n- `data/`：结构化后的工序、设备、距离数据。\n- `figures/`：各问题甘特图和总时长对比图。\n')
md.append('## 参考资料\n')
md.append('1. Ku & Beck, Mixed Integer Programming Models for Job Shop Scheduling: A Computational Analysis.\n2. Artigues, Lopez, Ayache, Schedule generation schemes for the job-shop problem with sequence-dependent setup times.\n3. Lan & Berkhout, PyJobShop: Solving scheduling problems with constraint programming in Python, 2025.\n4. SciPy Documentation, scipy.optimize.milp.\n')
(DOCS/'研究报告.md').write_text('\n'.join(md), encoding='utf-8')

# DOCX report
doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8); sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
styles = doc.styles
styles['Normal'].font.name = 'Noto Sans CJK SC'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
styles['Normal'].font.size = Pt(10.5)
for sty in ['Heading 1','Heading 2','Heading 3','Title']:
    styles[sty].font.name = 'Noto Sans CJK SC'
    styles[sty]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')

def add_para(text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix); run.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p

def set_cell_text(cell, text, bold=False, size=8):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.font.name = 'Noto Sans CJK SC'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    run.font.size = Pt(size)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(title, rows, cols, font_size=7.5):
    doc.add_heading(title, level=2)
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i,c in enumerate(cols):
        set_cell_text(hdr[i], c, True, font_size)
    for row in rows:
        cells = table.add_row().cells
        for i,c in enumerate(cols):
            set_cell_text(cells[i], row.get(c,''), False, font_size)
    return table

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('2026年五一数学建模B题\n多工序协同作业问题研究报告')
run.bold = True; run.font.size = Pt(20); run.font.name='Noto Sans CJK SC'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('研究报告、模型、代码与结果表汇总')
r.font.size = Pt(11)

doc.add_heading('摘要', level=1)
add_para('本报告围绕多个车间集中整修任务，建立了带有工序链约束、设备类型容量约束、跨车间序列相关转运时间的混合整数线性规划模型。题目本质上属于柔性作业车间调度与资源受限项目调度的交叉问题。对固定设备条件下的Q1、Q2、Q3，模型均由MILP求得全局最优解；对Q4，通过C车间关键工序链给出不可突破下界，并证明Q3已达到该下界，因此追加设备不再产生缩短工期的边际收益。')

doc.add_heading('一、主要结论', level=1)
add_table('表A 主要结果汇总', summary, ['问题','最短时长(s)','最短时长(HH:MM:SS)','求解状态'], font_size=8)
add_para('结论1：班组1独立完成A车间的最短时长为41600s，即11:33:20。')
add_para('结论2：仅使用班组1完成五个车间的最短时长为163764s，即45:29:24，主要受高速抛光机与自动传感多功能机单台瓶颈影响。')
add_para('结论3：使用两个班组后最短时长下降到123844s，即34:24:04。')
add_para('结论4：Q4在预算500000元约束下，若以最短完工时间为第一目标、采购成本为第二目标，则最优采购数量为0。原因是C车间固定工序链已构成全局关键路径，增购设备无法降低单工序处理时间。')

if (FIGS/'makespan_summary.png').exists():
    doc.add_picture(str(FIGS/'makespan_summary.png'), width=Inches(5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


doc.add_heading('二、问题分析与调研定位', level=1)
add_para('该问题同时包含三类典型调度结构：第一，车间内工序存在固定先后关系，属于项目调度中的precedence constraints；第二，同一类型设备有多台可替代机器，属于flexible job shop scheduling；第三，同一设备跨车间作业存在与前后车间相关的转运时间，可视为sequence-dependent setup time。文献中常用混合整数规划、约束规划与启发式算法处理此类问题。由于本题展开后只有27个工序、41个设备子任务，规模适合采用MILP获得可证明最优解。')
add_para('本报告采用MILP而非单纯贪心，是因为设备分配、设备排序、转运时间三者耦合，局部最早开工不一定给出全局最短完工时间。模型采用二元变量描述“设备分配”和“同设备任务先后顺序”，以Big-M线性化同机不重叠约束。')

doc.add_heading('三、数据预处理', level=1)
add_para('1. 工序展开：C3-C5重复三轮，因此C车间流程为C1、C2、C3_R1、C4_R1、C5_R1、C3_R2、C4_R2、C5_R2、C3_R3、C4_R3、C5_R3。')
add_para('2. 作业时间：对任一设备子任务，持续工作时间 = ceil(工程量 / 作业效率 * 3600)，单位为秒。')
add_para('3. 转运时间：设备速度均为2m/s，跨车间转运时间 = ceil(距离 / 2)。同一车间连续作业转运时间为0。')
add_para('4. 初始位置：设备初始位于对应班组驻地，首个车间作业前需考虑“班组-车间”距离。')

# Process duration overview small table
proc_rows=[]
for p in processes:
    dur_desc='；'.join([f"{ms.TYPE_ZH[t]}:{d}s" for t,d in p['durations'].items()])
    proc_rows.append({'车间':p['workshop'],'工序':p['pid'],'工程量':p['workload'],'设备作业时间':dur_desc,'工序完成下限(s)':p['maxdur']})
add_table('表B 展开后的工序与作业时间', proc_rows, ['车间','工序','工程量','设备作业时间','工序完成下限(s)'], font_size=7)


doc.add_heading('四、数学模型', level=1)
add_para('集合与参数：J为工序集合，O为设备子任务集合，M为设备集合。W_j表示工序j所属车间，tau_o表示子任务o所需设备类型，p_o表示子任务o处理时长，d_uv表示从地点u到地点v的转运时间。')
add_para('决策变量：S_j为工序j开始时间，T为最大完工时间；x_om为0-1变量，表示子任务o是否分配给设备m；y_abm为0-1变量，表示在同一设备m上子任务a是否排在子任务b之前。')
add_para('目标函数：min T。')
add_para('核心约束如下。')
add_para('(1) 每个设备子任务必须分配给一台且仅一台相应类型设备：sum_{m in M_tau(o)} x_om = 1。')
add_para('(2) 车间内部顺序约束：若工序j后继为k，则 S_k >= S_j + max_{o in O(j)} p_o。')
add_para('(3) 最大完工时间约束：T >= S_j + max_{o in O(j)} p_o。')
add_para('(4) 初始到达约束：若x_om=1，则S_{j(o)} >= d_{origin(m), W_j}。')
add_para('(5) 同机不重叠与转运约束：若a和b分配到同一台设备m，则必须满足 a在b前 或 b在a前。对应线性化形式为：S_b >= S_a + p_a + d_{W_a,W_b} 或 S_a >= S_b + p_b + d_{W_b,W_a}。')
add_para('由于所有持续时间和距离均确定，目标函数为最小化最大完工时间，固定设备版本可由MILP求得全局最优。')


doc.add_heading('五、求解过程与结果解释', level=1)
add_para('求解器使用 scipy.optimize.milp，其底层封装HiGHS混合整数线性规划求解器。每个固定设备问题先以T为目标求最优完工时间，再在T不变的条件下二次优化工序开始时间之和，使输出表格更紧凑，避免非关键工序被任意推迟。')
add_para('Q2相较Q1规模扩大后，单班组只有1台高速抛光机和1台自动传感多功能机，两个设备类型成为关键瓶颈。Q3引入班组2后，这两类瓶颈设备数量翻倍，工期降至C车间关键链下界。')


doc.add_heading('六、Q4采购策略分析', level=1)
add_para('C车间固定工序链不含初始转运的最短用时为：C1 10368s + C2 7406s + 3*(C3 6480s + C4 14400s + C5 14400s) = 123614s。两班组到C车间的最短初始转运时间为班组1到C，即460/2=230s。因此任何调度方案的总完工时间均不可能小于123844s。')
add_para('Q3已经达到123844s。因此在不改变设备效率、不允许拆分单一工序工作量、不改变C车间顺序的题设下，追加设备无法缩短总工期。Q4若以“最短完工时间”为唯一目标，存在多个等价采购方案；若增加“同等工期下采购成本最小”的理性次级目标，则采购0台为最优。')
add_table('表C Q4设备采购方案', table5, ['设备名称','班组1购买台数','班组2购买台数','设备单价(元)','小计(元)'], font_size=8)
add_para('购买设备总费用：0元。')


doc.add_heading('七、模型评价', level=1)
add_para('优点：模型完整刻画了设备分配、车间内工序顺序、设备跨车间转运、双设备工序同步开始等关键约束，并能给出可验证的全局最优解。')
add_para('局限：本模型默认双设备工序中较快设备完成自身工作后即可释放；若解释为两类设备必须共同占用至该工序整体结束，则需将该工序内所有设备子任务的占用时长统一改为该工序max时长。代码结构可直接修改p_o定义完成敏感性分析。')
add_para('可扩展方向：若题目规模扩大，可将MILP替换为CP-SAT、遗传算法、禁忌搜索或滚动时域启发式；若加入采购必须花完预算、设备效率差异或工序可拆分，则需扩展采购变量与加工模式变量。')


doc.add_heading('八、交付物索引', level=1)
add_para('1. docs/研究报告.docx：本文档。')
add_para('2. docs/研究报告.md：Markdown版本研究报告。')
add_para('3. code/model_solver.py：可复现实验代码，运行后生成结果CSV。')
add_para('4. data/：结构化输入数据。')
add_para('5. results/：表1至表5与汇总结果。')
add_para('6. figures/：Q1-Q4甘特图与总时长对比图。')


doc.add_heading('参考资料', level=1)
refs = [
    'Ku, W.-Y., Beck, J. C. Mixed Integer Programming Models for Job Shop Scheduling: A Computational Analysis.',
    'Artigues, C., Lopez, P., Ayache, P.-D. Schedule generation schemes for the job-shop problem with sequence-dependent setup times.',
    'Lan, L., Berkhout, J. PyJobShop: Solving scheduling problems with constraint programming in Python, 2025.',
    'SciPy Documentation: scipy.optimize.milp, Mixed-integer linear programming wrapper of HiGHS.',
]
for ref in refs:
    add_para(ref)

# Landscape appendix
new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
new_sec.orientation = WD_ORIENT.LANDSCAPE
new_sec.page_width, new_sec.page_height = new_sec.page_height, new_sec.page_width
new_sec.top_margin = Cm(1.2); new_sec.bottom_margin = Cm(1.2); new_sec.left_margin = Cm(1.2); new_sec.right_margin = Cm(1.2)

doc.add_heading('附录：可提交结果表', level=1)
add_table('表1 问题1结果', table1, ['序号','设备编号','起始时间','结束时间','持续工作时间(s)','工序编号'], font_size=7)
add_para('完成问题1任务的最短时长：41600(s)。')
add_table('表2 问题2结果', table2, ['序号','设备编号','起始时间','结束时间','持续工作时间(s)','工序编号'], font_size=6.5)
add_para('完成问题2任务的最短时长：163764(s)。')
add_table('表3 问题3结果', table3, ['序号','设备编号','起始时间','结束时间','持续工作时间(s)','工序编号','班组'], font_size=6.5)
add_para('完成问题3任务的最短时长：123844(s)。')
add_table('表4 问题4结果', table4, ['序号','设备编号','起始时间','结束时间','持续工作时间(s)','工序编号','班组'], font_size=6.5)
add_para('完成问题4任务的最短时长：123844(s)。')

out = DOCS/'研究报告.docx'
doc.save(out)
print(out)
